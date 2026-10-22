"""Exports Ashlie's quote workbook (7 tabs) to src/data/statementSheets.json + public/statement-logo.png.
Usage: python3 scripts/export_statements.py path/to/ASHLIES_QUOTES.xlsx"""
import sys, json, colorsys, re, zipfile, datetime
import openpyxl
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import get_column_letter

src = sys.argv[1]
wb = openpyxl.load_workbook(src)

# ---- theme colours ----
theme = []
try:
    xml = wb.loaded_theme.decode() if isinstance(wb.loaded_theme, bytes) else wb.loaded_theme
    order = ['lt1', 'dk1', 'lt2', 'dk2', 'accent1', 'accent2', 'accent3', 'accent4', 'accent5', 'accent6', 'hlink', 'folHlink']
    for name in order:
        m = re.search(r'<a:%s>(.*?)</a:%s>' % (name, name), xml, re.S)
        s = m.group(1) if m else ''
        r = re.search(r'(?:val|lastClr)="([0-9A-Fa-f]{6})"', s)
        theme.append(r.group(1).upper() if r else '000000')
except Exception:
    theme = ['FFFFFF', '000000', 'EEECE1', '1F497D', '4F81BD', 'C0504D', '9BBB59', '8064A2', '4BACC6', 'F79646', '0000FF', '800080']

def tint(hexc, t):
    r, g, b = [int(hexc[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + t) if t < 0 else l * (1 - t) + t
    r, g, b = colorsys.hls_to_rgb(h, min(max(l, 0), 1), s)
    return '%02X%02X%02X' % (round(r * 255), round(g * 255), round(b * 255))

def colour(c):
    if c is None: return None
    try:
        if c.type == 'rgb' and c.rgb and isinstance(c.rgb, str):
            return c.rgb[-6:].upper()
        if c.type == 'theme':
            base = theme[c.theme] if c.theme < len(theme) else '000000'
            return tint(base, c.tint or 0) if c.tint else base
        if c.type == 'indexed' and c.indexed is not None and c.indexed < len(COLOR_INDEX):
            return COLOR_INDEX[c.indexed][-6:].upper()
    except Exception:
        return None
    return None

def side(b):
    if b is None or not b.style: return None
    return [b.style, colour(b.color) or '000000']

def serial(d):
    if isinstance(d, datetime.datetime):
        return (d - datetime.datetime(1899, 12, 30)).total_seconds() / 86400
    return (datetime.datetime.combine(d, datetime.time()) - datetime.datetime(1899, 12, 30)).days

def px_col(w): return round((w if w else 8.43) * 7 + 5)

out = []
for ws in wb.worksheets:
    maxr, maxc = ws.max_row, ws.max_column
    # column widths (dimensions may cover ranges)
    cols = [px_col(ws.sheet_format.defaultColWidth or 8.43)] * maxc
    for k, d in ws.column_dimensions.items():
        lo, hi = d.min or 1, d.max or d.min or 1
        for c in range(lo, min(hi, maxc) + 1):
            cols[c - 1] = 0 if d.hidden else px_col(d.width)
    dh = ws.sheet_format.defaultRowHeight or 12.75
    rows = []
    for r in range(1, maxr + 1):
        d = ws.row_dimensions.get(r)
        h = d.height if d is not None and d.height else dh
        rows.append(0 if (d is not None and d.hidden) else round(h * 96 / 72))
    merges = [[m.min_row, m.min_col, m.max_row, m.max_col] for m in ws.merged_cells.ranges]
    mmap = {}
    for r1, c1, r2, c2 in merges: mmap[(r1, c1)] = (r2, c2)
    styles, sidx, cells = [], {}, {}
    def style_of(cell, r, c):
        f, a = cell.font, cell.alignment
        fill = cell.fill
        bg = colour(fill.fgColor) if fill and fill.fill_type == 'solid' else None
        bd = cell.border
        bt, bl, bb, br = side(bd.top), side(bd.left), side(bd.bottom), side(bd.right)
        if (r, c) in mmap:  # take right/bottom edges from the far corner cells of a merged block
            r2, c2 = mmap[(r, c)]
            br = side(ws.cell(r, c2).border.right) or br
            bb = side(ws.cell(r2, c).border.bottom) or bb
            bt = bt or side(ws.cell(r, c).border.top)
        s = {'ff': f.name, 'fs': f.sz, 'b': 1 if f.b else 0, 'i': 1 if f.i else 0, 'u': 1 if f.u else 0,
             'fc': colour(f.color), 'bg': bg, 'bt': bt, 'bl': bl, 'bb': bb, 'br': br,
             'ha': a.horizontal, 'va': a.vertical, 'wr': 1 if a.wrap_text else 0, 'rot': a.textRotation or 0, 'ind': a.indent or 0,
             'nf': cell.number_format}
        s = {k: v for k, v in s.items() if v not in (None, 0, False)}
        key = json.dumps(s, sort_keys=True)
        if key not in sidx:
            sidx[key] = len(styles); styles.append(s)
        return sidx[key]
    covered = set()
    for r1, c1, r2, c2 in merges:
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                if (r, c) != (r1, c1): covered.add((r, c))
    for r in range(1, maxr + 1):
        for c in range(1, maxc + 1):
            if (r, c) in covered: continue
            cell = ws.cell(r, c)
            v = cell.value
            if v is None and not cell.has_style: continue
            e = {'s': style_of(cell, r, c)}
            if isinstance(v, str) and v.startswith('='): e['f'] = v[1:]
            elif isinstance(v, (datetime.datetime, datetime.date)): e['v'] = serial(v)
            elif v is not None: e['v'] = v
            cells['%d,%d' % (r, c)] = e
    img = None
    if ws._images:
        a = ws._images[0].anchor
        try:
            img = {'from': [a._from.col, a._from.row, a._from.colOff // 9525, a._from.rowOff // 9525],
                   'to': [a.to.col, a.to.row, a.to.colOff // 9525, a.to.rowOff // 9525]}
        except Exception:
            img = None
    # which cells hold the person / decedent details (found next to their labels) so the app can fill them from the case
    fields = {}
    want = {'Name:': 'name', 'Relationship:': 'rel', 'Address:': 'addr', 'Phone:': 'phone', 'Decedent:': 'dec', 'Date of Death:': 'dod'}
    for k, e in cells.items():
        v = e.get('v')
        if isinstance(v, str) and v.strip() in want:
            r, c = map(int, k.split(','))
            end = mmap.get((r, c), (r, c))[1]
            for c2 in range(end + 1, maxc + 1):
                e2 = cells.get('%d,%d' % (r, c2))
                if e2 and ('v' in e2 or 'f' in e2 or styles[e2['s']].get('bb')):
                    fields[want[v.strip()]] = '%d,%d' % (r, c2); break
    ref = [f for f in out if f['fields']]
    if len(fields) < 6 and ref:   # same layout on every tab: fall back to the same offset from the label column
        labcol = {w: int(k.split(',')[1]) for k, e in cells.items() for w in [want.get(str(e.get('v')).strip())] if w}
        for key, off in (('name', 4), ('rel', 4), ('addr', 4), ('phone', 4), ('dec', 5), ('dod', 5)):
            if key not in fields and key in labcol:
                row = int(ref[0]['fields'][key].split(',')[0])
                fields[key] = '%d,%d' % (row, labcol[key] + off)
    # drop blank trailing columns / rows so the form scales up to fill a printed page
    def has_ink(e):
        st = styles[e['s']]
        return 'v' in e or 'f' in e or any(k in st for k in ('bt', 'bl', 'bb', 'br', 'bg'))
    usedc = max([int(k.split(',')[1]) for k, e in cells.items() if has_ink(e)] + [m[3] for m in merges] + [(img['to'][0] + 1) if img else 0])
    usedr = max([int(k.split(',')[0]) for k, e in cells.items() if has_ink(e)] + [m[2] for m in merges])
    cols, rows = cols[:usedc], rows[:usedr]
    cells = {k: e for k, e in cells.items() if int(k.split(',')[0]) <= usedr and int(k.split(',')[1]) <= usedc}
    merges = [[a, b, min(c, usedr), min(d, usedc)] for a, b, c, d in merges if a <= usedr and b <= usedc]
    out.append({'name': ws.title.strip(), 'fields': fields, 'cols': cols, 'rows': rows, 'merges': merges, 'styles': styles, 'cells': cells, 'image': img})

json.dump(out, open('src/data/statementSheets.json', 'w'), separators=(',', ':'))
z = zipfile.ZipFile(src)
media = [n for n in z.namelist() if n.startswith('xl/media/')]
if media:
    open('public/statement-logo.png', 'wb').write(z.read(media[0]))
print('sheets', [(s['name'], len(s['cells'])) for s in out], 'media', media[:1])
